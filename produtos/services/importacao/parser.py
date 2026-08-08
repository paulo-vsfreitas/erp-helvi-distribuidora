import csv
import io
import unicodedata
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


COLUNAS = {
    "categoria_comercial": "Categoria Comercial",
    "codigo_fornecedor": "Código do fornecedor",
    "codigo": "Código ERP",
    "modelo": "Modelo",
    "marca": "Marca",
    "colecao": "Coleção",
    "genero": "Gênero",
    "tipo_armacao": "Tipo de armação",
    "sem_variacao": "Sem variação",
    "cor_variacao": "Cor / Variação",
    "codigo_variacao": "Código da Variação",
    "custo": "Custo",
    "venda": "Venda",
    "estoque": "Estoque",
    "estoque_minimo": "Estoque mínimo",
    "observacoes": "Observações",
}


def _normalizar_texto(valor):
    if valor is None:
        return ""

    return str(valor).strip()


def _normalizar_cabecalho(valor):
    texto = _normalizar_texto(valor).lower()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(caractere)
    )

    return " ".join(texto.split())


CABECALHOS_NORMALIZADOS = {
    _normalizar_cabecalho(rotulo): chave
    for chave, rotulo in COLUNAS.items()
}


def _converter_decimal(valor, *, linha, campo, erros):
    if valor in (None, ""):
        return Decimal("0.00")

    if isinstance(valor, (int, float, Decimal)):
        try:
            return Decimal(str(valor))
        except InvalidOperation:
            pass

    texto = str(valor).strip()

    if not texto:
        return Decimal("0.00")

    texto = texto.replace("R$", "").strip()

    if "." in texto and "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")

    try:
        numero = Decimal(texto)
    except InvalidOperation:
        erros.append(
            f"Linha {linha}: {campo} inválido: {valor!r}."
        )
        return Decimal("0.00")

    if numero < 0:
        erros.append(
            f"Linha {linha}: {campo} não pode ser negativo."
        )

    return numero


def _converter_inteiro(valor, *, linha, campo, erros):
    if valor in (None, ""):
        return 0

    try:
        numero = int(valor)
    except (TypeError, ValueError):
        try:
            numero = int(Decimal(str(valor).replace(",", ".")))
        except (InvalidOperation, TypeError, ValueError):
            erros.append(
                f"Linha {linha}: {campo} inválido: {valor!r}."
            )
            return 0

    if numero < 0:
        erros.append(
            f"Linha {linha}: {campo} não pode ser negativo."
        )

    return numero


def _converter_booleano(valor):
    texto = _normalizar_cabecalho(valor)

    return texto in {
        "sim",
        "s",
        "1",
        "true",
        "verdadeiro",
        "yes",
    }


def _normalizar_categoria(valor, *, linha, erros):
    texto = _normalizar_cabecalho(valor)

    if texto in {"armacao", "armacoes"}:
        return "armacao"

    if texto in {"acessorio", "acessorios"}:
        return "acessorio"

    erros.append(
        f"Linha {linha}: Categoria Comercial inválida: "
        f"{valor!r}. Use Armação ou Acessório."
    )

    return ""


def _montar_linha(dados, numero_linha):
    erros = []

    categoria = _normalizar_categoria(
        dados.get("categoria_comercial"),
        linha=numero_linha,
        erros=erros,
    )

    sem_variacao = _converter_booleano(
        dados.get("sem_variacao")
    )

    codigo_fornecedor = _normalizar_texto(
        dados.get("codigo_fornecedor")
    ).upper()

    codigo = _normalizar_texto(
        dados.get("codigo")
    ).upper()

    modelo = _normalizar_texto(
        dados.get("modelo")
    )

    cor_variacao = _normalizar_texto(
        dados.get("cor_variacao")
    )

    codigo_variacao = _normalizar_texto(
        dados.get("codigo_variacao")
    ).upper()

    if not codigo_fornecedor and not codigo and not modelo:
        erros.append(
            f"Linha {numero_linha}: informe pelo menos "
            "Código do fornecedor, Código ERP ou Modelo."
        )

    if categoria == "acessorio":
        sem_variacao = True

    if categoria == "armacao" and not sem_variacao:
        if not codigo_variacao:
            erros.append(
                f"Linha {numero_linha}: informe o Código da "
                "Variação para a armação."
            )

    custo = _converter_decimal(
        dados.get("custo"),
        linha=numero_linha,
        campo="Custo",
        erros=erros,
    )

    venda = _converter_decimal(
        dados.get("venda"),
        linha=numero_linha,
        campo="Venda",
        erros=erros,
    )

    estoque = _converter_inteiro(
        dados.get("estoque"),
        linha=numero_linha,
        campo="Estoque",
        erros=erros,
    )

    estoque_minimo = _converter_inteiro(
        dados.get("estoque_minimo"),
        linha=numero_linha,
        campo="Estoque mínimo",
        erros=erros,
    )

    return {
        "linha": numero_linha,
        "categoria_comercial": categoria,
        "codigo_fornecedor": codigo_fornecedor,
        "codigo": codigo,
        "modelo": modelo,
        "marca": _normalizar_texto(
            dados.get("marca")
        ),
        "colecao": _normalizar_texto(
            dados.get("colecao")
        ),
        "genero": _normalizar_texto(
            dados.get("genero")
        ),
        "tipo_armacao": _normalizar_texto(
            dados.get("tipo_armacao")
        ),
        "sem_variacao": sem_variacao,
        "cor_variacao": cor_variacao,
        "codigo_variacao": codigo_variacao,
        "custo": custo,
        "venda": venda,
        "estoque": estoque,
        "estoque_minimo": estoque_minimo,
        "observacoes": _normalizar_texto(
            dados.get("observacoes")
        ),
        "erros": erros,
    }


def _mapear_cabecalhos(cabecalhos):
    mapa = {}

    for indice, cabecalho in enumerate(cabecalhos):
        normalizado = _normalizar_cabecalho(
            cabecalho
        )

        chave = CABECALHOS_NORMALIZADOS.get(
            normalizado
        )

        if chave:
            mapa[indice] = chave

    return mapa


def _validar_colunas(cabecalhos):
    encontrados = {
        _normalizar_cabecalho(valor)
        for valor in cabecalhos
        if valor not in (None, "")
    }

    faltantes = []

    for rotulo in COLUNAS.values():
        if _normalizar_cabecalho(rotulo) not in encontrados:
            faltantes.append(rotulo)

    return faltantes


def _ler_xlsx(arquivo):
    arquivo.seek(0)

    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=True,
    )

    if "Produtos" in workbook.sheetnames:
        worksheet = workbook["Produtos"]
    else:
        worksheet = workbook.active

    linhas = worksheet.iter_rows(values_only=True)

    try:
        cabecalhos = list(next(linhas))
    except StopIteration:
        return [], ["A planilha está vazia."]

    faltantes = _validar_colunas(cabecalhos)

    if faltantes:
        return [], [
            "Colunas obrigatórias ausentes: "
            + ", ".join(faltantes)
            + "."
        ]

    mapa = _mapear_cabecalhos(cabecalhos)

    resultado = []

    for numero_linha, valores in enumerate(
        linhas,
        start=2,
    ):
        if not any(
            valor not in (None, "")
            for valor in valores
        ):
            continue

        dados = {
            chave: (
                valores[indice]
                if indice < len(valores)
                else None
            )
            for indice, chave in mapa.items()
        }

        resultado.append(
            _montar_linha(
                dados,
                numero_linha,
            )
        )

    return resultado, []


def _ler_csv(arquivo):
    arquivo.seek(0)

    conteudo = arquivo.read()

    if isinstance(conteudo, bytes):
        try:
            conteudo = conteudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            conteudo = conteudo.decode(
                "latin-1"
            )

    amostra = conteudo[:4096]

    try:
        dialeto = csv.Sniffer().sniff(
            amostra,
            delimiters=";,",
        )
        delimitador = dialeto.delimiter
    except csv.Error:
        delimitador = ";"

    leitor = csv.reader(
        io.StringIO(conteudo),
        delimiter=delimitador,
    )

    try:
        cabecalhos = next(leitor)
    except StopIteration:
        return [], ["O arquivo CSV está vazio."]

    faltantes = _validar_colunas(cabecalhos)

    if faltantes:
        return [], [
            "Colunas obrigatórias ausentes: "
            + ", ".join(faltantes)
            + "."
        ]

    mapa = _mapear_cabecalhos(cabecalhos)

    resultado = []

    for numero_linha, valores in enumerate(
        leitor,
        start=2,
    ):
        if not any(
            _normalizar_texto(valor)
            for valor in valores
        ):
            continue

        dados = {
            chave: (
                valores[indice]
                if indice < len(valores)
                else None
            )
            for indice, chave in mapa.items()
        }

        resultado.append(
            _montar_linha(
                dados,
                numero_linha,
            )
        )

    return resultado, []


def ler_arquivo_produtos(arquivo):
    nome = arquivo.name.lower()

    if nome.endswith(".xlsx"):
        linhas, erros_arquivo = _ler_xlsx(
            arquivo
        )

    elif nome.endswith(".csv"):
        linhas, erros_arquivo = _ler_csv(
            arquivo
        )

    else:
        return {
            "linhas": [],
            "erros": [
                "Formato de arquivo não suportado."
            ],
        }

    erros = list(erros_arquivo)

    for linha in linhas:
        erros.extend(
            linha.get("erros", [])
        )

    return {
        "linhas": linhas,
        "erros": erros,
    }