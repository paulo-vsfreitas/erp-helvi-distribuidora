from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from catalogo.models import Colecao, Genero, Marca, TipoArmacao


COLUNAS_PRODUTOS = [
    ("Código do fornecedor", True),
    ("Código ERP", False),
    ("Modelo", False),
    ("Marca", False),
    ("Coleção", False),
    ("Gênero", False),
    ("Tipo de armação", False),
    ("Sem variação", False),
    ("Cor / Variação", False),
    ("Código da Variação", False),
    ("Custo", True),
    ("Venda", True),
    ("Estoque", True),
    ("Estoque mínimo", False),
    ("Observações", False),
]


COMENTARIOS = {
    "Código do fornecedor": (
        "Código utilizado pelo fornecedor para identificar o produto."
    ),
    "Código ERP": (
        "Opcional. Deixe vazio para novos produtos."
    ),
    "Modelo": (
        "Modelo ou descrição comercial do produto."
    ),
    "Marca": (
        "Informe uma marca existente no cadastro do ERP."
    ),
    "Coleção": (
        "Informe uma coleção existente no cadastro do ERP."
    ),
    "Gênero": (
        "Informe um gênero existente no cadastro do ERP. "
        "Ex.: Adulto, Infantil, Feminino, Masculino ou outro valor cadastrado."
    ),
    "Tipo de armação": (
        "Para armações, informe um tipo existente no cadastro do ERP. "
        "Ex.: Acetato ou Clipon."
    ),
    "Sem variação": (
        "Use SIM quando o produto não possuir variação. "
        "Caso contrário, use NAO."
    ),
    "Cor / Variação": (
        "Cor da armação ou descrição da variação. "
        "Ex.: Preto, Tartaruga, Azul."
    ),
    "Código da Variação": (
        "Identificador da variação dentro do produto. "
        "Ex.: C1, C2, C3. Para Clipon, pode ser 5 em 1."
    ),
    "Custo": (
        "Preço de custo do produto."
    ),
    "Venda": (
        "Preço de venda do produto."
    ),
    "Estoque": (
        "Quantidade física disponível. "
        "Um Clipon 5 em 1 com cinco lentes continua sendo estoque 1."
    ),
    "Estoque mínimo": (
        "Opcional. Quantidade mínima desejada em estoque."
    ),
    "Observações": (
        "Informações adicionais sobre o produto."
    ),
}


def gerar_modelo_importacao_produtos():
    workbook = Workbook()

    planilha = workbook.active
    planilha.title = "Produtos"

    preenchimento_obrigatorio = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    preenchimento_opcional = PatternFill(
        fill_type="solid",
        fgColor="FFFFFF",
    )

    for indice, (nome, obrigatorio) in enumerate(
        COLUNAS_PRODUTOS,
        start=1,
    ):
        celula = planilha.cell(
            row=1,
            column=indice,
            value=nome,
        )

        celula.font = Font(bold=True)

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        celula.fill = (
            preenchimento_obrigatorio
            if obrigatorio
            else preenchimento_opcional
        )

        comentario = COMENTARIOS.get(nome)

        if comentario:
            celula.comment = Comment(
                comentario,
                "ERP Helvi",
            )

        largura = max(len(nome) + 4, 18)

        planilha.column_dimensions[
            get_column_letter(indice)
        ].width = largura

    planilha.freeze_panes = "A2"

    planilha.auto_filter.ref = (
        f"A1:{get_column_letter(len(COLUNAS_PRODUTOS))}1"
    )

    # Exemplo de armação comum
    planilha.append(
        [
            "ROMA-001",
            "",
            "Roma",
            "",
            "",
            "Adulto",
            "Acetato",
            "NAO",
            "Preto",
            "C1",
            35.00,
            89.90,
            1,
            0,
            "",
        ]
    )

    # Exemplo de Clipon 5 em 1
    planilha.append(
        [
            "CLIP-001",
            "",
            "Clipon 5 em 1",
            "",
            "",
            "Adulto",
            "Clipon",
            "NAO",
            "Preto",
            "5 em 1",
            40.00,
            119.90,
            1,
            0,
            "Acompanha 5 lentes clip-on.",
        ]
    )

    referencias = workbook.create_sheet("Referencias")

    referencias.append(
        [
            "Marcas",
            "Coleções",
            "Gêneros",
            "Tipos de armação",
        ]
    )

    for celula in referencias[1]:
        celula.font = Font(bold=True)

    marcas = list(
        Marca.objects.filter(ativo=True)
        .order_by("nome")
        .values_list("nome", flat=True)
    )

    colecoes = list(
        Colecao.objects.filter(ativo=True)
        .order_by("nome")
        .values_list("nome", flat=True)
    )

    generos = list(
        Genero.objects.filter(ativo=True)
        .order_by("nome")
        .values_list("nome", flat=True)
    )

    tipos_armacao = list(
        TipoArmacao.objects.filter(ativo=True)
        .order_by("nome")
        .values_list("nome", flat=True)
    )

    maior_lista = max(
        len(marcas),
        len(colecoes),
        len(generos),
        len(tipos_armacao),
        1,
    )

    for indice in range(maior_lista):
        referencias.append(
            [
                marcas[indice] if indice < len(marcas) else "",
                colecoes[indice] if indice < len(colecoes) else "",
                generos[indice] if indice < len(generos) else "",
                tipos_armacao[indice] if indice < len(tipos_armacao) else "",
            ]
        )

    for indice in range(1, 6):
        referencias.column_dimensions[
            get_column_letter(indice)
        ].width = 24

    buffer = BytesIO()

    workbook.save(buffer)
    buffer.seek(0)

    return buffer